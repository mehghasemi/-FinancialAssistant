import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

import app.database as database
from scripts.import_workbook import run_import


class WorkbookImportTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_data_dir, self.old_database_path = database.DATA_DIR, database.DATABASE_PATH
        database.DATA_DIR = Path(self.temp_dir.name)
        database.DATABASE_PATH = Path(self.temp_dir.name) / "financial-assistant-test.db"

    def tearDown(self):
        database.DATA_DIR, database.DATABASE_PATH = self.old_data_dir, self.old_database_path
        self.temp_dir.cleanup()

    def test_archives_every_sheet_and_maps_budget_and_installment(self):
        workbook_path = Path(self.temp_dir.name) / "source.xlsx"
        workbook = Workbook()
        obligations = workbook.active
        obligations.title = "تعهدات مالی"
        obligations.cell(6, 2, "نوع")
        obligations.cell(6, 3, "مبلغ کل")
        obligations.cell(6, 4, "تاریخ")
        obligations.cell(6, 5, "تاریخ سررسید")
        obligations.cell(6, 7, "وضعیت")
        obligations.cell(6, 8, "مبلغ")
        obligations.cell(7, 2, "وام نمونه")
        obligations.cell(7, 3, 2_000_000)
        obligations.cell(7, 4, "1405/07/01")
        obligations.cell(7, 5, datetime(2026, 9, 23))
        obligations.cell(7, 7, "پرداخت شده")
        obligations.cell(7, 8, 1_000_000)
        budget = workbook.create_sheet("بودجه 1405")
        budget.cell(2, 1, "سرفصل بودجه")
        budget.cell(2, 2, "فروردین")
        budget.cell(3, 1, "خانه")
        budget.cell(3, 2, 5_000_000)
        other = workbook.create_sheet("داده خام")
        other.cell(1, 1, "نگهداری شود")
        workbook.save(workbook_path)

        result = run_import(workbook_path, create_backup=False)

        self.assertFalse(result["already_imported"])
        self.assertEqual(result["installments"], 1)
        self.assertEqual(result["payments"], 1)
        self.assertEqual(result["budget_items"], 1)
        with database.connection() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM imported_rows").fetchone()[0], 5)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM commitments").fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM installments").fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM payments").fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM budget_items").fetchone()[0], 1)

        second = run_import(workbook_path, create_backup=False)
        self.assertTrue(second["already_imported"])


if __name__ == "__main__":
    unittest.main()
